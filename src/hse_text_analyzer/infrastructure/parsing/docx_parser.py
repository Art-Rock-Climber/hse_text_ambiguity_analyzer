from io import BytesIO

from hse_text_analyzer.application.ports.document_parser import DocumentParser, UploadedDocument
from hse_text_analyzer.domain.enums import SourceFileType
from hse_text_analyzer.domain.exceptions import DocumentParsingError
from hse_text_analyzer.domain.models.document import Document, DocumentMetadata, DocumentSection


class DocxDocumentParser(DocumentParser):
    name = "docx_parser"
    supported_file_types = frozenset({SourceFileType.DOCX})

    def parse(self, uploaded_document: UploadedDocument) -> Document:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise DocumentParsingError("python-docx is required to parse DOCX files.") from exc

        try:
            docx_document = DocxDocument(BytesIO(uploaded_document.content))
        except Exception as exc:
            message = f"Could not parse DOCX file '{uploaded_document.file_name}'."
            raise DocumentParsingError(message) from exc

        paragraphs = [paragraph.text.strip() for paragraph in docx_document.paragraphs]
        text = "\n\n".join(paragraph for paragraph in paragraphs if paragraph)
        sections = self._build_sections(text)
        return Document(
            title=uploaded_document.file_name,
            raw_text=text,
            metadata=DocumentMetadata(
                file_name=uploaded_document.file_name,
                file_type=SourceFileType.DOCX,
                document_type=uploaded_document.declared_document_type,
                parser_name=self.name,
            ),
            sections=sections,
        )

    def _build_sections(self, text: str) -> list[DocumentSection]:
        return [
            DocumentSection(title=None, text=text, start_char=0, end_char=len(text))
        ] if text else []
